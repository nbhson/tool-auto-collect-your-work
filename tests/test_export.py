"""Unit tests for /api/export endpoints."""

from datetime import UTC, datetime
from io import BytesIO

from openpyxl import load_workbook


class TestExportExcel:
    """Tests for GET /api/export/excel."""

    def _create_log(self, client, **kwargs):
        defaults = {
            "title": "Test log",
            "activity_type": "meeting",
            "time_spent_minutes": 30,
        }
        defaults.update(kwargs)
        return client.post("/api/logs", json=defaults).json()

    def test_export_empty(self, client):
        """Export of empty DB returns valid Excel with 2 sheets."""
        resp = client.get("/api/export/excel")
        assert resp.status_code == 200
        assert (
            resp.headers["content-type"]
            == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert "attachment" in resp.headers["content-disposition"]
        assert "worklog_export" in resp.headers["content-disposition"]

        # Verify it's a valid xlsx
        wb = load_workbook(BytesIO(resp.content))
        assert "Work Logs" in wb.sheetnames
        assert "Summary" in wb.sheetnames

    def test_export_with_data(self, client):
        """Export includes created entries."""
        self._create_log(client, title="Meeting one", source="manual")
        self._create_log(client, title="Commit two", source="github")

        resp = client.get("/api/export/excel")
        assert resp.status_code == 200

        wb = load_workbook(BytesIO(resp.content))
        ws = wb["Work Logs"]

        # Row 1 is headers, rows 2+ are data
        assert ws.cell(1, 1).value == "ID"  # header
        # Find the data rows (skip header)
        data_rows = [row for row in ws.iter_rows(min_row=2) if row[0].value]
        titles = [row[3].value for row in data_rows]  # Title is column 4
        assert "Meeting one" in titles
        assert "Commit two" in titles

    def test_export_summary_sheet(self, client, db_session):
        """Summary sheet includes counts per source."""
        from database.models import WorkLog

        db_session.add(
            WorkLog(
                source="github",
                activity_type="commit",
                title="GH log",
                activity_timestamp=datetime.now(UTC).replace(tzinfo=None),
            )
        )
        db_session.add(
            WorkLog(
                source="manual",
                activity_type="meeting",
                title="Manual log",
                activity_timestamp=datetime.now(UTC).replace(tzinfo=None),
            )
        )
        db_session.commit()

        resp = client.get("/api/export/excel")
        wb = load_workbook(BytesIO(resp.content))
        ws2 = wb["Summary"]

        metrics = {row[0].value: row[1].value for row in ws2.iter_rows(min_row=2)}

        assert metrics["Total Logs"] == 2
        assert metrics["GitHub Activities"] == 1
        assert metrics["Manual Entries"] == 1

    def test_export_by_source(self, client, db_session):
        """Export by source filters correctly."""
        from database.models import WorkLog

        db_session.add(
            WorkLog(
                source="github",
                activity_type="commit",
                title="GH commit",
                activity_timestamp=datetime.now(UTC).replace(tzinfo=None),
            )
        )
        db_session.add(
            WorkLog(
                source="manual",
                activity_type="meeting",
                title="Manual entry",
                activity_timestamp=datetime.now(UTC).replace(tzinfo=None),
            )
        )
        db_session.commit()

        resp = client.get("/api/export/excel/github")
        assert resp.status_code == 200

        wb = load_workbook(BytesIO(resp.content))
        ws = wb["Work Logs"]

        data_rows = [row for row in ws.iter_rows(min_row=2) if row[0].value]
        sources = {row[1].value for row in data_rows}
        assert sources == {"github"}

    def test_export_invalid_source(self, client):
        """Invalid source returns 422."""
        resp = client.get("/api/export/excel/invalid_source")
        assert resp.status_code == 422
        assert resp.json()["error"] == "ValidationError"
